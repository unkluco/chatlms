from __future__ import annotations

import csv
import ctypes
import hashlib
import json
import os
import random
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

from health_monitor import (
    cleanup_old_attachments,
    find_browser_memory_mb,
    get_process_memory_mb,
    write_health,
)

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.json"
STATE_PATH = ROOT / "processed.json"
STATE_BACKUP_PATH = ROOT / "processed.json.bak"
PENDING_PATH = ROOT / "pending_comments.json"
PENDING_BACKUP_PATH = ROOT / "pending_comments.json.bak"
SESSION_USER_PATH = ROOT / "session_username.txt"
IDENTITY_PATH = ROOT / "account_identities.json"
PROFILE_DIR = ROOT / ".browser_profile"
ATTACHMENT_DIR = ROOT / ".attachments"
RUNTIME_DIR = ROOT / ".bot_runtime"
PID_PATH = RUNTIME_DIR / "bot.pid"
LOG_DIR = ROOT / "logs"
LOG_FILE = LOG_DIR / "bot.log"
MAX_PROCESSED_PER_ACCOUNT = 50000
MAX_PENDING_COMMENTS_PER_ACCOUNT = 200
LOG_RETENTION_DAYS = 30
MAX_LOG_ARCHIVES = 30

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".py", ".java", ".c", ".cc", ".cpp", ".cxx",
    ".h", ".hpp", ".cs", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css",
    ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".conf", ".log", ".sql",
    ".sh", ".bat", ".ps1", ".go", ".rs", ".php", ".rb", ".kt", ".kts", ".swift"
}
SUPPORTED_ATTACHMENT_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".csv"}

def cleanup_old_logs(retention_days=LOG_RETENTION_DAYS, max_archives=MAX_LOG_ARCHIVES):
    try:
        if not LOG_DIR.exists():
            return 0
        cutoff = time.time() - max(1, retention_days) * 86400
        archives = sorted(
            LOG_DIR.glob("bot-*.log"),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        removed = 0
        for index, path in enumerate(archives):
            try:
                if index >= max_archives or path.stat().st_mtime < cutoff:
                    path.unlink(missing_ok=True)
                    removed += 1
            except Exception:
                continue
        return removed
    except Exception:
        return 0


def log(msg):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    print(line, flush=True)
    try:
        LOG_DIR.mkdir(exist_ok=True)
        if LOG_FILE.exists() and LOG_FILE.stat().st_size > 10 * 1024 * 1024:
            backup = LOG_DIR / f"bot-{datetime.now().strftime('%Y%m%d-%H%M%S')}.log"
            LOG_FILE.replace(backup)
            cleanup_old_logs()
        with LOG_FILE.open("a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass

def validate_config(cfg):
    if not isinstance(cfg, dict):
        raise RuntimeError("config.json phai la mot JSON object.")

    errors = []
    base_url = str(cfg.get("lms_base_url", "")).strip()
    parsed = urlparse(base_url)
    if not base_url:
        errors.append("thieu lms_base_url")
    elif parsed.scheme not in ("https", "http") or not parsed.netloc:
        errors.append("lms_base_url khong hop le")
    elif parsed.scheme != "https" and not bool(cfg.get("allow_insecure_http", False)):
        errors.append("lms_base_url phai dung HTTPS (hoac allow_insecure_http=true)")

    if not str(cfg.get("username", "")).strip():
        errors.append("thieu username")
    if not str(cfg.get("password", "")).strip():
        errors.append("thieu password")

    provider = str(cfg.get("provider", "groq")).strip().lower() or "groq"
    if provider != "groq":
        errors.append(f"provider chua ho tro: {provider}")
    if provider == "groq" and not str(cfg.get("groq_api_key", "")).strip():
        errors.append("thieu groq_api_key")

    numeric_rules = [
        ("poll_interval_seconds", 2),
        ("max_posts_per_scan", 1),
        ("fast_scan_pages", 1),
        ("deep_scan_interval_seconds", 30),
        ("deep_scan_max_posts", 1),
        ("deep_scan_max_pages", 2),
        ("max_answer_length", 100),
        ("retry_count", 0),
        ("retry_delay_seconds", 0),
        ("retry_max_delay_seconds", 0),
        ("cooldown_seconds", 0),
        ("ignored_recheck_seconds", 30),
        ("reconnect_delay_seconds", 2),
        ("login_retry_seconds", 2),
        ("health_interval_seconds", 30),
        ("maintenance_interval_seconds", 300),
        ("browser_restart_hours", 1),
        ("max_python_memory_mb", 256),
        ("max_browser_children_memory_mb", 512),
        ("watchdog_stale_seconds", 120),
        ("watchdog_check_seconds", 5),
    ]
    for key, minimum in numeric_rules:
        if key not in cfg:
            continue
        try:
            value = float(cfg[key])
            if value < minimum:
                errors.append(f"{key} phai >= {minimum}")
        except (TypeError, ValueError):
            errors.append(f"{key} phai la so")

    attachments = cfg.get("attachments", {})
    if attachments is not None and not isinstance(attachments, dict):
        errors.append("attachments phai la object")
    elif isinstance(attachments, dict):
        attachment_numeric = [
            ("max_file_size_mb", 1),
            ("max_files_per_post", 1),
            ("max_extracted_chars_per_file", 100),
            ("max_total_attachment_chars", 100),
        ]
        for key, minimum in attachment_numeric:
            if key not in attachments:
                continue
            try:
                value = float(attachments[key])
                if value < minimum:
                    errors.append(
                        f"attachments.{key} phai >= {minimum}"
                    )
            except (TypeError, ValueError):
                errors.append(f"attachments.{key} phai la so")

    trigger_position = str(
        cfg.get("trigger_position", "content")
    ).strip().lower()
    if trigger_position not in ("content", "title", "both"):
        errors.append("trigger_position phai la content/title/both")

    if errors:
        raise RuntimeError("Config khong hop le: " + "; ".join(errors))
    return cfg


def load_config():
    if not CONFIG_PATH.exists():
        raise RuntimeError(
            "Khong tim thay config.json. Sao chep config.example.json thanh config.json."
        )
    try:
        cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            f"config.json khong phai JSON hop le: dong {exc.lineno}, cot {exc.colno}"
        ) from exc

    env_overrides = {
        "LMS_USERNAME": "username",
        "LMS_PASSWORD": "password",
        "GROQ_API_KEY": "groq_api_key",
    }
    for env_name, key in env_overrides.items():
        value = os.getenv(env_name)
        if value:
            cfg[key] = value

    return validate_config(cfg)

def debug_log(cfg, msg):
    if bool(cfg.get("debug", False)):
        log("[DEBUG] " + msg)

def acquire_single_instance(cfg):
    if os.name != "nt":
        return None

    # Mot project chi dung mot persistent browser profile, nen chi cho phep
    # mot bot chay trong cung thu muc du config dang la tai khoan nao.
    identity = (
        str(ROOT).strip().lower()
        + "|"
        + str(cfg.get("lms_base_url", "")).strip().lower()
    )
    suffix = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]
    mutex_name = "Local\\IUH_LMS_BLOG_BOT_" + suffix

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p]
    kernel32.CreateMutexW.restype = ctypes.c_void_p
    kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
    kernel32.CloseHandle.restype = ctypes.c_bool

    handle = kernel32.CreateMutexW(None, False, mutex_name)
    if not handle:
        raise RuntimeError("Khong tao duoc single-instance mutex.")

    ERROR_ALREADY_EXISTS = 183
    if ctypes.get_last_error() == ERROR_ALREADY_EXISTS:
        kernel32.CloseHandle(handle)
        return False

    return kernel32, handle

def release_single_instance(lock):
    if not lock or lock is False:
        return
    try:
        kernel32, handle = lock
        kernel32.CloseHandle(handle)
    except Exception:
        pass


def create_runtime_pid():
    """Luu PID de phat hien zombie process. Mutex van la lop chinh."""
    try:
        RUNTIME_DIR.mkdir(exist_ok=True)
        PID_PATH.write_text(str(os.getpid()), encoding="utf-8")
    except Exception:
        pass


def clear_runtime_pid():
    try:
        PID_PATH.unlink(missing_ok=True)
    except Exception:
        pass


def load_runtime_pid():
    try:
        return int(PID_PATH.read_text(encoding="utf-8").strip())
    except Exception:
        return None

def _atomic_write_json(path, data, backup_path=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    if backup_path is not None and path.exists():
        try:
            json.loads(path.read_text(encoding="utf-8"))
            shutil.copy2(path, backup_path)
        except Exception:
            pass

    tmp_path = path.with_name(path.name + ".tmp")
    with tmp_path.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(tmp_path, path)


def _read_json_with_backup(path, backup_path=None, default=None):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as primary_exc:
        if backup_path is not None and backup_path.exists():
            try:
                data = json.loads(backup_path.read_text(encoding="utf-8"))
                log(f"Canh bao: {path.name} bi loi; da khoi phuc tu {backup_path.name}.")
                _atomic_write_json(path, data)
                return data
            except Exception:
                pass
        raise RuntimeError(
            f"state_corrupt:{path.name}:{type(primary_exc).__name__}"
        ) from primary_exc


def load_processed(account_key="default"):
    data = _read_json_with_backup(
        STATE_PATH,
        STATE_BACKUP_PATH,
        default=None,
    )
    if data is None:
        return set()
    if not isinstance(data, dict):
        raise RuntimeError("state_corrupt:processed.json:not_object")

    if "accounts" in data:
        accounts = data.get("accounts", {})
        if not isinstance(accounts, dict):
            raise RuntimeError("state_corrupt:processed.json:accounts")
        return set(map(str, accounts.get(account_key, [])))
    return set(map(str, data.get("processed_entry_ids", [])))


def _trim_processed_ids(items, limit=MAX_PROCESSED_PER_ACCOUNT):
    values = list({str(value) for value in items})
    if len(values) <= limit:
        return sorted(
            values,
            key=lambda value: (0, int(value)) if value.isdigit() else (1, value),
        )

    numeric = [value for value in values if value.isdigit()]
    other = [value for value in values if not value.isdigit()]
    numeric.sort(key=int)
    other.sort()
    combined = numeric + other
    return combined[-limit:]


def save_processed(items, account_key="default"):
    old = _read_json_with_backup(
        STATE_PATH,
        STATE_BACKUP_PATH,
        default=None,
    )
    data = old if isinstance(old, dict) else {"version": 3, "accounts": {}, "account_meta": {}}

    if "accounts" not in data:
        legacy = data.get("processed_entry_ids", [])
        data = {
            "version": 3,
            "accounts": {"legacy": list(map(str, legacy))},
            "account_meta": {},
        }

    data["version"] = 3
    data.setdefault("account_meta", {})
    accounts = data.setdefault("accounts", {})
    if not isinstance(accounts, dict):
        raise RuntimeError("state_corrupt:processed.json:accounts")

    for key, values in list(accounts.items()):
        accounts[key] = _trim_processed_ids(values)
    accounts[account_key] = _trim_processed_ids(items)

    _atomic_write_json(STATE_PATH, data, STATE_BACKUP_PATH)


def _load_pending_data():
    data = _read_json_with_backup(
        PENDING_PATH,
        PENDING_BACKUP_PATH,
        default=None,
    )
    if data is None:
        return {"version": 1, "accounts": {}}
    if not isinstance(data, dict):
        raise RuntimeError("state_corrupt:pending_comments.json:not_object")
    accounts = data.setdefault("accounts", {})
    if not isinstance(accounts, dict):
        raise RuntimeError("state_corrupt:pending_comments.json:accounts")
    data["version"] = 1
    return data


def get_pending_comment(account_key, eid):
    data = _load_pending_data()
    item = data.get("accounts", {}).get(str(account_key), {}).get(str(eid))
    return item if isinstance(item, dict) else None


def save_pending_comment(account_key, eid, answer):
    data = _load_pending_data()
    accounts = data.setdefault("accounts", {})
    pending = accounts.setdefault(str(account_key), {})
    pending[str(eid)] = {
        "answer": str(answer),
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }

    if len(pending) > MAX_PENDING_COMMENTS_PER_ACCOUNT:
        ordered = sorted(
            pending.items(),
            key=lambda pair: str(pair[1].get("created_at", "")),
        )
        for old_eid, _ in ordered[: len(pending) - MAX_PENDING_COMMENTS_PER_ACCOUNT]:
            pending.pop(old_eid, None)

    _atomic_write_json(PENDING_PATH, data, PENDING_BACKUP_PATH)


def clear_pending_comment(account_key, eid):
    if not PENDING_PATH.exists():
        return
    data = _load_pending_data()
    accounts = data.setdefault("accounts", {})
    pending = accounts.get(str(account_key))
    if isinstance(pending, dict):
        pending.pop(str(eid), None)
        if not pending:
            accounts.pop(str(account_key), None)
    _atomic_write_json(PENDING_PATH, data, PENDING_BACKUP_PATH)


def prune_pending_comments(account_key, processed_ids):
    if not PENDING_PATH.exists():
        return 0
    data = _load_pending_data()
    accounts = data.setdefault("accounts", {})
    pending = accounts.get(str(account_key))
    if not isinstance(pending, dict):
        return 0

    processed_ids = {str(value) for value in processed_ids}
    removed = 0
    for eid in list(pending):
        if str(eid) in processed_ids:
            pending.pop(eid, None)
            removed += 1

    if not pending:
        accounts.pop(str(account_key), None)
    if removed:
        _atomic_write_json(PENDING_PATH, data, PENDING_BACKUP_PATH)
    return removed


def get_account_meta(account_key):
    data = _read_json_with_backup(
        STATE_PATH,
        STATE_BACKUP_PATH,
        default=None,
    )
    if not isinstance(data, dict):
        return {}
    meta = data.get("account_meta", {}).get(str(account_key), {})
    return dict(meta) if isinstance(meta, dict) else {}


def update_account_meta(account_key, **fields):
    data = _read_json_with_backup(
        STATE_PATH,
        STATE_BACKUP_PATH,
        default=None,
    )
    if data is None:
        data = {"version": 3, "accounts": {}, "account_meta": {}}
    if not isinstance(data, dict):
        raise RuntimeError("state_corrupt:processed.json:not_object")

    data["version"] = max(3, int(data.get("version", 0) or 0))
    data.setdefault("accounts", {})
    account_meta = data.setdefault("account_meta", {})
    if not isinstance(account_meta, dict):
        raise RuntimeError("state_corrupt:processed.json:account_meta")
    meta = account_meta.setdefault(str(account_key), {})
    if not isinstance(meta, dict):
        meta = {}
        account_meta[str(account_key)] = meta
    for key, value in fields.items():
        if value is None:
            meta.pop(str(key), None)
        else:
            meta[str(key)] = value
    _atomic_write_json(STATE_PATH, data, STATE_BACKUP_PATH)


def _entry_id_int(value):
    text = str(value or "").strip()
    return int(text) if text.isdigit() else None

def entry_id(url):
    try:
        x = parse_qs(urlparse(url).query).get("entryid")
        return x[0] if x else None
    except Exception:
        return None

def launch_browser(pw, headless=False, ignore_https_errors=False):
    for channel in ("chrome", "msedge"):
        try:
            log(f"Mo trinh duyet {channel}...")
            return pw.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE_DIR),
                channel=channel,
                headless=headless,
                ignore_https_errors=ignore_https_errors,
                viewport={"width": 1400, "height": 900},
                args=["--disable-notifications"],
            )
        except Exception:
            pass
    return pw.chromium.launch_persistent_context(
        user_data_dir=str(PROFILE_DIR),
        headless=headless,
        ignore_https_errors=ignore_https_errors,
        viewport={"width": 1400, "height": 900},
    )

def logged_out(page):
    # Khong dua vao URL /login/: Moodle van cho phep nguoi da dang nhap
    # mo /login/index.php va hien trang "ban da dang nhap" ma khong co form.
    # Chi xem la da logout khi form dang nhap thuc su xuat hien.
    return (
        page.locator('input[name="username"]').count() > 0
        and page.locator('input[name="password"]').count() > 0
    )

def browser_needs_restart(page, exc=None):
    try:
        if page.is_closed():
            return True
    except Exception:
        return True

    message = str(exc or "").lower()
    dead_markers = (
        "target page, context or browser has been closed",
        "browser has been closed",
        "context has been closed",
        "page has been closed",
        "connection closed",
    )
    return any(marker in message for marker in dead_markers)

def _normalize_account_name(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def load_account_identities():
    if not IDENTITY_PATH.exists():
        return {}
    try:
        data = json.loads(IDENTITY_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_account_identity(username, user_id, display_name):
    data = load_account_identities()
    data[str(username)] = {
        "userid": str(user_id),
        "display_name": str(display_name or "").strip(),
    }
    IDENTITY_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def discover_session_identity(page, cfg):
    base_url = str(cfg.get("lms_base_url", "https://lms.iuh.edu.vn")).rstrip("/")
    page.goto(base_url + "/", wait_until="domcontentloaded", timeout=45000)

    name = ""
    name_selectors = [
        '[data-region="usermenu"] .usertext',
        '.usermenu .usertext',
        '#user-menu-toggle .usertext',
        '.logininfo a[href*="/user/profile.php?id="]',
    ]
    for selector in name_selectors:
        try:
            loc = page.locator(selector)
            if loc.count() > 0:
                value = loc.first.inner_text(timeout=2000).strip()
                if value:
                    name = value
                    break
        except Exception:
            pass

    profile_selectors = [
        '[data-region="usermenu"] a[href*="/user/profile.php?id="]',
        '.usermenu a[href*="/user/profile.php?id="]',
        '.logininfo a[href*="/user/profile.php?id="]',
    ]
    for selector in profile_selectors:
        try:
            links = page.locator(selector)
            for i in range(min(links.count(), 20)):
                href = links.nth(i).get_attribute("href") or ""
                values = parse_qs(urlparse(href).query).get("id")
                if values and values[0].isdigit():
                    if not name:
                        try:
                            name = links.nth(i).inner_text(timeout=1500).strip()
                        except Exception:
                            pass
                    debug_log(cfg, f"Session profile: userid={values[0]}, name={name!r}")
                    return values[0], name
        except Exception:
            pass

    return None, name


def _login_with_config(page, cfg, reason=""):
    base_url = str(cfg.get("lms_base_url", "https://lms.iuh.edu.vn")).rstrip("/")
    login_url = base_url + "/login/index.php"
    username = str(cfg.get("username", "")).strip()
    password = str(cfg.get("password", ""))
    if not username or not password:
        raise RuntimeError("Thieu username/password trong config.json.")

    if reason:
        log(reason)
    # Profile nay chi danh cho bot, xoa cookie la cach chac chan nhat de cat session cu.
    try:
        page.context.clear_cookies()
    except Exception:
        pass

    page.goto(login_url, wait_until="domcontentloaded", timeout=45000)
    if not logged_out(page):
        raise RuntimeError("Khong dua LMS ve duoc man hinh dang nhap sau khi xoa cookie.")

    log(f"Dang dang nhap LMS theo username trong config: {username}")
    page.locator('input[name="username"]').fill(username, timeout=10000)
    page.locator('input[name="password"]').fill(password, timeout=10000)
    page.locator('#loginbtn, button[type="submit"], input[type="submit"]').first.click(timeout=10000)
    try:
        page.wait_for_load_state("domcontentloaded", timeout=30000)
    except Exception:
        pass
    page.wait_for_timeout(1000)

    if logged_out(page):
        raise RuntimeError("Dang nhap theo config that bai; kiem tra username/password hoac CAPTCHA/MFA.")

    user_id, display_name = discover_session_identity(page, cfg)
    if not user_id:
        raise RuntimeError("Dang nhap thanh cong nhung khong xac dinh duoc userid cua tai khoan.")

    save_account_identity(username, user_id, display_name)
    SESSION_USER_PATH.write_text(username, encoding="utf-8")
    log(f"Da xac minh tai khoan config: {display_name or '(khong doc duoc ten)'} | userid={user_id}")
    return user_id, display_name


def ensure_login(page, cfg):
    base_url = str(cfg.get("lms_base_url", "https://lms.iuh.edu.vn")).rstrip("/")
    login_url = base_url + "/login/index.php"
    retry_seconds = max(2, int(cfg.get("login_retry_seconds", 5)))
    username = str(cfg.get("username", "")).strip()

    while True:
        try:
            identities = load_account_identities()
            expected = identities.get(username)
            old_user = SESSION_USER_PATH.read_text(encoding="utf-8").strip() if SESSION_USER_PATH.exists() else ""

            page.goto(login_url, wait_until="domcontentloaded", timeout=45000)
            if logged_out(page):
                return _login_with_config(page, cfg, "LMS chua dang nhap; dung tai khoan trong config.json.")

            current_id, current_name = discover_session_identity(page, cfg)

            # Tai khoan nay chua tung duoc xac minh bang credential trong config.
            # Khong tin session persistent hien tai: dang nhap lai mot lan de lap moc chuan.
            if not expected:
                return _login_with_config(
                    page, cfg,
                    f"Chua co moc danh tinh cho username {username}; dang nhap lai de xac minh dung tai khoan."
                )

            expected_id = str(expected.get("userid", ""))
            expected_name = str(expected.get("display_name", "")).strip()
            mismatch = False
            reasons = []

            if old_user and old_user != username:
                mismatch = True
                reasons.append(f"config doi {old_user} -> {username}")
            if not current_id or str(current_id) != expected_id:
                mismatch = True
                reasons.append(f"userid session={current_id}, expected={expected_id}")
            if expected_name and current_name and _normalize_account_name(current_name) != _normalize_account_name(expected_name):
                mismatch = True
                reasons.append(f"ten session={current_name!r}, expected={expected_name!r}")

            if mismatch:
                return _login_with_config(
                    page, cfg,
                    "Phat hien session LMS khong dung tai khoan config (" + "; ".join(reasons) + "). Dang xuat session cu va dang nhap lai."
                )

            SESSION_USER_PATH.write_text(username, encoding="utf-8")
            log(f"Session dung tai khoan config: {current_name or expected_name or username} | userid={current_id}")
            return current_id, current_name or expected_name

        except KeyboardInterrupt:
            raise
        except Exception as exc:
            if browser_needs_restart(page, exc):
                raise RuntimeError("browser_restart_required") from exc
            log(f"Khong xac minh/dang nhap duoc LMS: {type(exc).__name__}: {exc}")
            log(f"Cho {retry_seconds}s roi thu lai...")
            time.sleep(retry_seconds)


def discover_user_id(page, cfg):
    user_id, _ = discover_session_identity(page, cfg)
    if user_id:
        return user_id
    raise RuntimeError("Khong xac dinh duoc userid cua session dang dang nhap.")

def _next_blog_page_url(page, current_url):
    selectors = [
        'a[rel="next"]',
        '.pagination a[aria-label*="Next" i]',
        '.pagination a[aria-label*="Sau" i]',
        '.pagination a[aria-label*="Tiếp" i]',
        '.pagination a:has-text("Next")',
        '.pagination a:has-text("Sau")',
        '.pagination a:has-text("Tiếp")',
    ]
    for selector in selectors:
        try:
            loc = page.locator(selector)
            if loc.count() > 0:
                href = loc.first.get_attribute("href")
                if href and href != current_url:
                    return href
        except Exception:
            pass

    try:
        current_page = parse_qs(urlparse(current_url).query).get("page", ["-1"])[0]
        current_page = int(current_page) if str(current_page).lstrip("-").isdigit() else -1
        candidates = page.locator('.pagination a[href*="page="], nav a[href*="page="]')
        best = None
        for i in range(min(candidates.count(), 100)):
            href = candidates.nth(i).get_attribute("href") or ""
            page_value = parse_qs(urlparse(href).query).get("page", [None])[0]
            if page_value is None or not str(page_value).isdigit():
                continue
            page_number = int(page_value)
            if page_number > current_page and (best is None or page_number < best[0]):
                best = (page_number, href)
        return best[1] if best else None
    except Exception:
        return None


def scan_blog_entries(
    page,
    blog_url,
    limit,
    skip_ids=None,
    max_pages=1,
    ignore_through_entry_id=None,
):
    """Quet blog co pagination, chi dem bai chua xu ly vao limit."""
    skip_ids = {str(value) for value in (skip_ids or set())}
    max_pages = max(1, int(max_pages or 1))
    limit = None if limit is None else max(1, int(limit))
    floor = _entry_id_int(ignore_through_entry_id)

    out = []
    seen_entries = set()
    visited_pages = set()
    current_url = blog_url
    pages_scanned = 0

    while current_url and pages_scanned < max_pages:
        if current_url in visited_pages:
            return out, {
                "pages_scanned": pages_scanned,
                "exhausted": True,
                "hit_limit": False,
                "hit_floor": False,
            }
        visited_pages.add(current_url)

        page.goto(current_url, wait_until="domcontentloaded", timeout=45000)
        if logged_out(page):
            raise RuntimeError("login_required")

        hrefs = page.locator('a[href*="entryid="]').evaluate_all(
            "els => els.map(e => e.href).filter(Boolean)"
        )
        page_ids = []
        for href in hrefs:
            eid = entry_id(href)
            if not eid or eid in seen_entries:
                continue
            seen_entries.add(eid)
            page_ids.append(eid)

            eid_int = _entry_id_int(eid)
            if floor is not None and eid_int is not None and eid_int <= floor:
                continue
            if eid in skip_ids:
                continue

            out.append((eid, href))
            if limit is not None and len(out) >= limit:
                return out, {
                    "pages_scanned": pages_scanned + 1,
                    "exhausted": False,
                    "hit_limit": True,
                    "hit_floor": False,
                }

        pages_scanned += 1

        numeric_page_ids = [
            value
            for value in (_entry_id_int(eid) for eid in page_ids)
            if value is not None
        ]
        if floor is not None and numeric_page_ids and max(numeric_page_ids) <= floor:
            return out, {
                "pages_scanned": pages_scanned,
                "exhausted": True,
                "hit_limit": False,
                "hit_floor": True,
            }

        next_url = _next_blog_page_url(page, current_url)
        if not next_url:
            return out, {
                "pages_scanned": pages_scanned,
                "exhausted": True,
                "hit_limit": False,
                "hit_floor": False,
            }
        current_url = next_url

    return out, {
        "pages_scanned": pages_scanned,
        "exhausted": False,
        "hit_limit": False,
        "hit_floor": False,
    }


def get_entries(page, blog_url, limit):
    entries, _ = scan_blog_entries(
        page,
        blog_url,
        limit,
        max_pages=1,
    )
    return entries

def first_text(page, selectors):
    for selector in selectors:
        try:
            loc = page.locator(selector)
            if loc.count() > 0:
                text = loc.first.inner_text(timeout=3000).strip()
                if text:
                    return text
        except Exception:
            pass
    return ""

def extract_entry_texts(page):
    title = first_text(page, [
        ".blog_entry .subject",
        ".blog-entry .subject",
        '[data-region="blog-post"] .subject',
        ".page-header-headings h1",
        "article h1",
        "article h2",
        "h1",
    ])
    content = first_text(page, [
        ".blog_entry .content",
        ".blog-entry .content",
        ".blog_entry .no-overflow",
        ".blog-entry .no-overflow",
        '[data-region="blog-post"] .content',
        "article .content",
        "article .no-overflow",
        "#region-main .blog_entry",
        "#region-main article",
        "#region-main",
    ])
    return title, content

def _safe_filename(url, fallback="attachment"):
    name = Path(unquote(urlparse(url).path)).name or fallback
    name = re.sub(r'[<>:"/\\|?*]', "_", name).strip(" .")
    return name or fallback

def _attachment_ext(url):
    return Path(unquote(urlparse(url).path)).suffix.lower()

def extract_attachment_links(page, cfg):
    settings = cfg.get("attachments", {})
    if not isinstance(settings, dict) or not settings.get("enabled", False):
        return []

    max_files = max(1, int(settings.get("max_files_per_post", 3)))
    base_host = urlparse(str(cfg.get("lms_base_url", ""))).netloc.lower()
    selectors = [
        '.blog_entry a[href]',
        '.blog-entry a[href]',
        '[data-region="blog-post"] a[href]',
        'article a[href]',
        '#region-main a[href*="/pluginfile.php/"]',
        '#region-main a[href*="/draftfile.php/"]',
    ]

    found = []
    seen = set()
    for selector in selectors:
        try:
            links = page.locator(selector)
            for i in range(min(links.count(), 100)):
                href = links.nth(i).evaluate("e => e.href") or ""
                if not href or href in seen:
                    continue
                parsed = urlparse(href)
                if parsed.netloc and base_host and parsed.netloc.lower() != base_host:
                    continue
                if "/pluginfile.php/" not in parsed.path and "/draftfile.php/" not in parsed.path:
                    continue
                ext = _attachment_ext(href)
                if ext not in SUPPORTED_ATTACHMENT_EXTENSIONS:
                    continue
                seen.add(href)
                text = (links.nth(i).inner_text(timeout=1000) or "").strip()
                found.append({
                    "url": href,
                    "name": text or _safe_filename(href),
                    "ext": ext,
                })
                if len(found) >= max_files:
                    return found
        except Exception:
            pass
    return found

def match_command(page, cfg):
    title, content = extract_entry_texts(page)

    detect_pattern = str(cfg.get("detect_pattern", "#bot")).strip()
    if not detect_pattern:
        return None

    default_system_prompt = str(cfg.get("default_system_prompt", "")).strip()
    commands = cfg.get("commands", {})
    if not isinstance(commands, dict):
        commands = {}

    position = str(cfg.get("trigger_position", "content")).strip().lower()
    if position not in ("content", "title", "both"):
        position = "content"

    sources = []
    if position in ("content", "both"):
        sources.append(("content", content))
    if position in ("title", "both"):
        sources.append(("title", title))

    for source_name, source_text in sources:
        detect_pos = source_text.find(detect_pattern)
        if detect_pos < 0:
            continue

        if source_name == "content":
            prompt = source_text[detect_pos + len(detect_pattern):].strip()
        else:
            title_after = source_text[detect_pos + len(detect_pattern):].strip()
            prompt = (title_after + "\n" + content).strip()

        selected_command = None
        selected_system_prompt = default_system_prompt

        for tag in sorted(commands.keys(), key=len, reverse=True):
            tag = str(tag)
            if not tag:
                continue
            tag_pos = prompt.find(tag)
            if tag_pos < 0:
                continue
            command_cfg = commands.get(tag) or {}
            selected_command = tag
            selected_system_prompt = str(command_cfg.get("system_prompt", "")).strip() or default_system_prompt
            prompt = (prompt[:tag_pos] + prompt[tag_pos + len(tag):]).strip()
            break

        return detect_pattern, selected_command, selected_system_prompt, prompt

    return None

def download_attachment(page, eid, item, cfg):
    settings = cfg.get("attachments", {}) or {}
    max_size = max(1, int(settings.get("max_file_size_mb", 15))) * 1024 * 1024

    try:
        response = page.context.request.get(item["url"], timeout=45000)
    except Exception as exc:
        return None, f"download loi: {type(exc).__name__}: {exc}"

    if not response.ok:
        return None, f"HTTP {response.status}"

    headers = response.headers or {}
    content_type = str(headers.get("content-type", "")).lower()
    if "text/html" in content_type and item.get("ext") not in (".html", ".htm"):
        raise RuntimeError("login_required")

    try:
        content_length = int(headers.get("content-length", "0") or 0)
    except Exception:
        content_length = 0

    if content_length and content_length > max_size:
        return None, f"file vuot gioi han {settings.get('max_file_size_mb', 15)} MB"

    body = response.body()
    if len(body) > max_size:
        return None, f"file vuot gioi han {settings.get('max_file_size_mb', 15)} MB"

    entry_dir = ATTACHMENT_DIR / str(eid)
    entry_dir.mkdir(parents=True, exist_ok=True)

    filename = _safe_filename(item["url"], item.get("name") or "attachment")
    path = entry_dir / filename

    if path.exists():
        stem, suffix = path.stem, path.suffix
        n = 2
        while path.exists():
            path = entry_dir / f"{stem}_{n}{suffix}"
            n += 1

    path.write_bytes(body)
    return path, None

def _decode_text_file(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "utf-8", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")

def _extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"[Trang {index}]\n{text}")
    return "\n\n".join(parts)

def _extract_docx(path):
    from docx import Document

    doc = Document(str(path))
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            parts.append(f"[Bang {table_index}]\n" + "\n".join(rows))

    return "\n\n".join(parts)

def _extract_xlsx(path, max_chars):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    total = 0

    try:
        for ws in wb.worksheets:
            header = f"[Sheet: {ws.title}]"
            parts.append(header)
            total += len(header) + 1

            for row in ws.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if not any(values):
                    continue
                line = " | ".join(values)
                parts.append(line)
                total += len(line) + 1
                if total >= max_chars:
                    parts.append("[Da cat bot vi vuot gioi han trich xuat]")
                    return "\n".join(parts)
    finally:
        wb.close()

    return "\n".join(parts)

def _extract_csv(path, max_chars):
    text = _decode_text_file(path)
    rows = []
    total = 0

    try:
        reader = csv.reader(text.splitlines())
        for row in reader:
            line = " | ".join(str(value) for value in row)
            rows.append(line)
            total += len(line) + 1
            if total >= max_chars:
                rows.append("[Da cat bot vi vuot gioi han trich xuat]")
                break
        return "\n".join(rows)
    except Exception:
        return text[:max_chars]

def extract_attachment_text(path, cfg):
    settings = cfg.get("attachments", {}) or {}
    max_chars = max(1000, int(settings.get("max_extracted_chars_per_file", 20000)))
    ext = path.suffix.lower()

    if ext in TEXT_EXTENSIONS:
        text = _decode_text_file(path)
    elif ext == ".pdf":
        text = _extract_pdf(path)
        if not text.strip():
            return "", "PDF khong co text trich xuat duoc; OCR/anh chua duoc ho tro"
    elif ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".xlsx":
        text = _extract_xlsx(path, max_chars)
    elif ext == ".csv":
        text = _extract_csv(path, max_chars)
    else:
        return "", f"chua ho tro dinh dang {ext}"

    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n[Da cat bot vi vuot gioi han trich xuat]"
    return text, None

def collect_attachment_context(page, eid, cfg):
    settings = cfg.get("attachments", {}) or {}
    if not settings.get("enabled", False):
        return ""

    links = extract_attachment_links(page, cfg)
    if not links:
        return ""

    max_total = max(1000, int(settings.get("max_total_attachment_chars", 40000)))
    delete_after = bool(settings.get("delete_after_processing", True))
    entry_dir = ATTACHMENT_DIR / str(eid)

    sections = []
    used_chars = 0

    try:
        log(f"Entry {eid}: tim thay {len(links)} file dinh kem ho tro.")
        for item in links:
            path, error = download_attachment(page, eid, item, cfg)
            if error:
                log(f"Entry {eid}: bo qua {item.get('name')}: {error}")
                continue

            try:
                text, error = extract_attachment_text(path, cfg)
            except Exception as exc:
                log(f"Entry {eid}: loi doc {path.name}: {type(exc).__name__}: {exc}")
                continue

            if error:
                log(f"Entry {eid}: bo qua {path.name}: {error}")
                continue
            if not text:
                continue

            remaining = max_total - used_chars
            if remaining <= 0:
                break
            text = text[:remaining]
            section = f"===== FILE: {path.name} =====\n{text}"
            sections.append(section)
            used_chars += len(text)

            log(f"Entry {eid}: da doc {path.name} ({len(text)} ky tu).")
            if used_chars >= max_total:
                log(f"Entry {eid}: da dat gioi han tong noi dung file {max_total} ky tu.")
                break
    finally:
        if delete_after and entry_dir.exists():
            shutil.rmtree(entry_dir, ignore_errors=True)

    return "\n\n".join(sections)

def call_groq(prompt, system_prompt, cfg):
    api_key = str(cfg.get("groq_api_key", "")).strip()
    if not api_key:
        raise RuntimeError("Thieu groq_api_key trong config.json")

    model = str(cfg.get("ai_model", "")).strip() or "openai/gpt-oss-120b"
    max_chars = max(100, int(cfg.get("max_answer_length", 6000)))
    max_tokens = max(128, min(4096, max_chars // 2 + 128))

    try:
        from groq import (
            AuthenticationError,
            Groq,
            NotFoundError,
            PermissionDeniedError,
        )
        client = Groq(api_key=api_key, timeout=90.0)
        response = client.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": system_prompt or "Tra loi dung trong tam va ro rang."
                },
                {"role": "user", "content": prompt}
            ],
            temperature=0.3,
            max_completion_tokens=max_tokens,
            reasoning_effort="low",
            include_reasoning=False
        )
    except (AuthenticationError, PermissionDeniedError, NotFoundError) as exc:
        raise RuntimeError(
            f"ai_permanent:{type(exc).__name__}:{exc}"
        ) from exc
    except Exception as exc:
        raise RuntimeError(f"Groq SDK error: {type(exc).__name__}: {exc}") from exc

    answer = str(response.choices[0].message.content or "").strip()
    if not answer:
        raise RuntimeError("Groq tra ve noi dung rong")
    return answer[:max_chars]

def call_ai(prompt, system_prompt, cfg):
    provider = str(cfg.get("provider", "groq")).strip().lower() or "groq"
    retries = max(0, int(cfg.get("retry_count", 2)))
    retry_delay = max(0.0, float(cfg.get("retry_delay_seconds", 2)))
    retry_max_delay = max(
        retry_delay,
        float(cfg.get("retry_max_delay_seconds", 30)),
    )
    last_error = None

    for attempt in range(retries + 1):
        try:
            if provider == "groq":
                return call_groq(prompt, system_prompt, cfg)
            raise RuntimeError(f"Provider chua duoc ho tro: {provider}")
        except Exception as exc:
            if str(exc).startswith("ai_permanent:"):
                raise
            last_error = exc
            debug_log(
                cfg,
                f"AI attempt {attempt + 1} loi: {type(exc).__name__}: {exc}",
            )
            if attempt < retries and retry_delay > 0:
                base_delay = min(retry_max_delay, retry_delay * (2 ** attempt))
                jitter = random.uniform(0, min(1.0, base_delay * 0.2))
                wait_seconds = base_delay + jitter
                debug_log(cfg, f"AI retry sau {wait_seconds:.1f}s.")
                time.sleep(wait_seconds)

    raise RuntimeError(str(last_error))

def _normalize_comment_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _ensure_comments_open(page):
    try:
        loaded_messages = page.locator(
            "ul.comment-list li .comment-message .no-overflow, "
            "ul.comment-list li .comment-message .text_to_html"
        )
        if loaded_messages.count() > 0:
            return
    except Exception:
        pass

    toggles = page.locator("a.comment-link")
    if toggles.count() > 0:
        try:
            toggles.last.click(timeout=5000)
            page.wait_for_timeout(800)
        except Exception:
            pass


def comment_exists(page, reply, expected_user_id=None):
    target = _normalize_comment_text(reply)
    if not target:
        return False

    _ensure_comments_open(page)
    items = page.locator("ul.comment-list li")
    for index in range(items.count()):
        item = items.nth(index)
        try:
            text_locator = item.locator(".no-overflow, .text_to_html").first
            text = _normalize_comment_text(text_locator.inner_text(timeout=1500))
        except Exception:
            continue
        if text != target:
            continue

        if expected_user_id:
            try:
                author = item.locator(
                    '.user a[href*="/user/"], .picture a[href*="/user/"]'
                ).first
                href = author.get_attribute("href") or ""
                author_id = parse_qs(urlparse(href).query).get("id", [None])[0]
                if str(author_id or "") != str(expected_user_id):
                    continue
            except Exception:
                continue
        return True
    return False


def _comment_textarea(page):
    selectors = [
        'textarea[name="content"]',
        'textarea[id^="dlg-content-"]',
        '.comment-area textarea',
    ]
    for selector in selectors:
        loc = page.locator(selector)
        if loc.count() > 0:
            return loc.last
    return None


def _comment_save_control(page):
    selectors = [
        'a[id^="comment-action-post-"]',
        '.comment-area a:has-text("Lưu lời bình")',
        '.comment-area a:has-text("Save comment")',
        '.comment-area button:has-text("Lưu lời bình")',
        '.comment-area button:has-text("Save comment")',
    ]
    for selector in selectors:
        loc = page.locator(selector)
        if loc.count() > 0:
            return loc.last
    return None


def post_comment(page, reply, expected_user_id=None):
    # Idempotency check: neu request truoc da duoc LMS nhan nhung bot crash
    # truoc khi ghi processed.json, khong post lai.
    if comment_exists(page, reply, expected_user_id):
        return True, "Binh luan da ton tai tren LMS; xac nhan lai thanh cong."

    textarea = _comment_textarea(page)
    if textarea is None or not textarea.is_visible():
        _ensure_comments_open(page)
        textarea = _comment_textarea(page)

    if textarea is None:
        return False, "Khong tim thay o binh luan."

    try:
        textarea.wait_for(state="visible", timeout=10000)
    except PlaywrightTimeoutError:
        return False, "Khong tim thay o binh luan."

    textarea.fill(reply)
    save_btn = _comment_save_control(page)
    if save_btn is None:
        return False, "Khong tim thay nut luu binh luan."

    save_btn.click(timeout=10000)

    deadline = time.time() + 12
    while time.time() < deadline:
        page.wait_for_timeout(500)
        if comment_exists(page, reply, expected_user_id):
            return True, "Da gui va xac minh binh luan tren LMS."

    # Neu AJAX khong cap nhat DOM, reload mot lan de xac minh server da nhan.
    try:
        current_url = page.url
        page.goto(current_url, wait_until="domcontentloaded", timeout=45000)
        if logged_out(page):
            raise RuntimeError("login_required")
        if comment_exists(page, reply, expected_user_id):
            return True, "Da gui; xac minh thanh cong sau khi tai lai trang."
    except RuntimeError:
        raise
    except Exception:
        pass

    return False, "Da bam gui nhung chua xac minh duoc binh luan; se thu lai an toan."


def process_entry(page, eid, url, cfg, account_key, user_id):
    page.goto(url, wait_until="domcontentloaded", timeout=45000)
    if logged_out(page):
        raise RuntimeError("login_required")

    pending = get_pending_comment(account_key, eid)
    answer = ""
    if pending and str(pending.get("answer", "")).strip():
        answer = str(pending["answer"]).strip()
        log(f"Entry {eid}: tim thay cau tra loi pending; kiem tra LMS truoc khi gui lai.")
    else:
        matched = match_command(page, cfg)
        if not matched:
            return "ignored"

        detect_pattern, command_tag, system_prompt, prompt = matched
        if not prompt:
            log(
                f"Entry {eid}: co mau kich hoat {detect_pattern!r} "
                "nhung khong co noi dung cau hoi; bo qua."
            )
            return "ignored"

        attachment_context = collect_attachment_context(page, eid, cfg)
        if attachment_context:
            prompt = (
                prompt
                + "\n\n===== NOI DUNG FILE DINH KEM =====\n"
                + attachment_context
            )

        mode = command_tag if command_tag else "default"
        log(f"Phat hien {detect_pattern!r} o entry {eid}; mode={mode}. Dang goi AI...")
        answer = call_ai(prompt, system_prompt, cfg)
        save_pending_comment(account_key, eid, answer)
        log(
            f"Entry {eid}: AI tra loi {len(answer)} ky tu; "
            "da luu pending truoc khi gui binh luan."
        )

    ok, detail = post_comment(page, answer, expected_user_id=user_id)
    log(f"Entry {eid}: {detail}")
    if ok:
        # Chua xoa pending o day. Main phai save processed truoc roi moi xoa
        # pending de khong co cua so crash gay comment trung.
        return "completed"
    return "retry"

def ensure_account_scan_floor(page, blog_url, account_key, processed, only_new):
    if not only_new:
        return None

    meta = get_account_meta(account_key)
    floor = _entry_id_int(meta.get("ignore_through_entry_id"))
    if floor is not None:
        return str(floor)

    processed_numeric = [
        value for value in (_entry_id_int(eid) for eid in processed)
        if value is not None
    ]
    if processed_numeric:
        floor = max(processed_numeric)
        update_account_meta(
            account_key,
            ignore_through_entry_id=str(floor),
            initialized_at=datetime.now().isoformat(timespec="seconds"),
            source="migrated_from_processed",
        )
        log(
            f"only_new_posts=true: tao moc migration entryid<={floor} "
            f"cho tai khoan {account_key}."
        )
        return str(floor)

    current, _ = scan_blog_entries(
        page,
        blog_url,
        limit=100,
        skip_ids=None,
        max_pages=1,
    )
    numeric = [
        value for value in (_entry_id_int(eid) for eid, _ in current)
        if value is not None
    ]
    floor = max(numeric) if numeric else 0
    update_account_meta(
        account_key,
        ignore_through_entry_id=str(floor),
        initialized_at=datetime.now().isoformat(timespec="seconds"),
        source="initial_baseline",
    )
    log(
        f"only_new_posts=true: tao moc entryid<={floor}; "
        "cac bai ton tai truoc moc se khong duoc xu ly."
    )
    return str(floor)


def main():
    cfg = load_config()

    instance_lock = acquire_single_instance(cfg)
    if instance_lock is False:
        log("Bot cho project/profile LMS nay dang chay o mot process khac.")
        log("Khong khoi dong instance thu hai de tranh binh luan trung.")
        return 2

    create_runtime_pid()

    processed = set()
    in_progress = set()

    base_url = str(cfg["lms_base_url"]).rstrip("/")
    commands = cfg.get("commands", {})
    only_new = bool(cfg.get("only_new_posts", False))
    interval = max(2, int(cfg.get("poll_interval_seconds", 5)))
    limit = max(1, int(cfg.get("max_posts_per_scan", 10)))
    fast_scan_pages = max(1, int(cfg.get("fast_scan_pages", 1)))
    deep_scan_interval = max(30, int(cfg.get("deep_scan_interval_seconds", 300)))
    deep_scan_max_posts = max(limit, int(cfg.get("deep_scan_max_posts", 500)))
    deep_scan_max_pages = max(2, int(cfg.get("deep_scan_max_pages", 100)))
    cooldown = max(0.0, float(cfg.get("cooldown_seconds", 0)))
    ignored_recheck_seconds = max(
        30,
        int(cfg.get("ignored_recheck_seconds", 300)),
    )
    health_interval = max(30, int(cfg.get("health_interval_seconds", 60)))
    maintenance_interval = max(300, int(cfg.get("maintenance_interval_seconds", 3600)))
    browser_restart_hours = max(1, float(cfg.get("browser_restart_hours", 6)))
    max_python_memory_mb = max(256, int(cfg.get("max_python_memory_mb", 800)))
    max_browser_children_memory_mb = max(512, int(cfg.get("max_browser_children_memory_mb", 1500)))
    reconnect_delay = max(2, int(cfg.get("reconnect_delay_seconds", 5)))
    provider = str(cfg.get("provider", "groq") or "groq")
    model = str(cfg.get("ai_model", "")).strip() or "openai/gpt-oss-120b"

    cleanup_old_logs()
    log("=== IUH LMS BLOG BOT ===")
    write_health("starting")
    names = ", ".join(commands.keys()) if isinstance(commands, dict) else "(none)"
    detect_pattern = str(cfg.get("detect_pattern", "#bot"))
    log(f"Mau kich hoat: {detect_pattern!r}; scan moi {interval}s")
    log(f"Lenh phu: {names}")
    log(f"AI provider: {provider}; model: {model}")
    log(f"Tu dong reconnect sau {reconnect_delay}s neu mat ket noi.")

    attachment_cfg = cfg.get("attachments", {}) or {}
    if attachment_cfg.get("enabled", False):
        log(
            "File dinh kem: ON | "
            f"toi da {attachment_cfg.get('max_files_per_post', 3)} file/bai | "
            f"{attachment_cfg.get('max_file_size_mb', 15)} MB/file"
        )
    else:
        log("File dinh kem: OFF")

    scan_floor_by_account = {}
    last_deep_scan_by_account = {}
    ignored_recent = {}
    last_health = 0
    last_maintenance = 0
    browser_started_at = time.time()

    with sync_playwright() as pw:
        while True:
            context = None
            try:
                context = launch_browser(
                    pw,
                    bool(cfg.get("headless", False)),
                    bool(cfg.get("ignore_https_errors", False)),
                )
                browser_started_at = time.time()
                write_health("browser_started")
                page = context.pages[0] if context.pages else context.new_page()

                user_id, display_name = ensure_login(page, cfg)
                expected_username = str(cfg.get("username", "")).strip()
                account_key = f"{expected_username}:{user_id}"
                log(
                    f"Tai khoan da xac minh: username={expected_username}, "
                    f"ten={display_name or '(khong doc duoc)'}, userid={user_id}"
                )
                blog_url = f"{base_url}/blog/index.php?userid={user_id}"
                processed = load_processed(account_key)
                stale_pending = prune_pending_comments(account_key, processed)
                if stale_pending:
                    debug_log(
                        cfg,
                        f"Da don {stale_pending} pending cu da co trong processed state.",
                    )
                log("Da xac dinh blog cua dung tai khoan config.")
                write_health("running", account=account_key, memory_mb=get_process_memory_mb(), browser_children_mb=find_browser_memory_mb(os.getpid()))

                scan_floor = ensure_account_scan_floor(
                    page,
                    blog_url,
                    account_key,
                    processed,
                    only_new,
                )
                scan_floor_by_account[account_key] = scan_floor

                while True:
                    try:
                        now = time.time()
                        removed_temp_files = 0
                        removed_log_files = 0

                        if now - last_maintenance >= maintenance_interval:
                            removed_temp_files = cleanup_old_attachments(ATTACHMENT_DIR, 24)
                            removed_log_files = cleanup_old_logs()
                            last_maintenance = now

                        if now - last_health >= health_interval:
                            python_memory_mb = get_process_memory_mb()
                            browser_children_mb = find_browser_memory_mb(os.getpid())
                            write_health(
                                "running",
                                account=account_key,
                                memory_mb=python_memory_mb,
                                browser_children_mb=browser_children_mb,
                                removed_temp_files=removed_temp_files,
                                removed_log_files=removed_log_files,
                            )
                            last_health = now
                            if (
                                (python_memory_mb or 0) > max_python_memory_mb
                                or (browser_children_mb or 0) > max_browser_children_memory_mb
                            ):
                                log(
                                    "Phat hien bo nho bot vuot nguong "
                                    f"(Python={python_memory_mb}MB, children={browser_children_mb}MB). "
                                    "Yeu cau supervisor khoi dong lai toan bo process."
                                )
                                raise RuntimeError("process_restart_required")

                        if now - browser_started_at >= browser_restart_hours * 3600:
                            log("Browser da chay du chu ky. Restart de giai phong tai nguyen.")
                            raise RuntimeError("browser_restart_required")

                        last_deep = last_deep_scan_by_account.get(account_key, 0)
                        do_deep_scan = now - last_deep >= deep_scan_interval
                        scan_limit = deep_scan_max_posts if do_deep_scan else limit
                        scan_pages = deep_scan_max_pages if do_deep_scan else fast_scan_pages
                        expired_ignored = [
                            eid
                            for eid, retry_at in ignored_recent.items()
                            if retry_at <= now
                        ]
                        for eid in expired_ignored:
                            ignored_recent.pop(eid, None)
                        recent_ignored_ids = set(ignored_recent)

                        entries, scan_meta = scan_blog_entries(
                            page,
                            blog_url,
                            scan_limit,
                            skip_ids=processed | in_progress | recent_ignored_ids,
                            max_pages=scan_pages,
                            ignore_through_entry_id=scan_floor_by_account.get(account_key),
                        )
                        if do_deep_scan:
                            last_deep_scan_by_account[account_key] = now
                            debug_log(
                                cfg,
                                "Deep scan: "
                                f"{scan_meta['pages_scanned']} trang, "
                                f"{len(entries)} bai chua xu ly, "
                                f"exhausted={scan_meta['exhausted']}, "
                                f"hit_limit={scan_meta['hit_limit']}."
                            )

                        for eid, url in entries:
                            if eid in processed or eid in in_progress:
                                continue

                            in_progress.add(eid)
                            try:
                                status = process_entry(
                                    page,
                                    eid,
                                    url,
                                    cfg,
                                    account_key,
                                    user_id,
                                )
                                if status == "completed":
                                    processed.add(eid)
                                    # Thu tu durability rat quan trong:
                                    # 1) save processed 2) moi clear pending.
                                    save_processed(processed, account_key)
                                    clear_pending_comment(account_key, eid)
                                    ignored_recent.pop(eid, None)
                                    if cooldown > 0:
                                        time.sleep(cooldown)
                                elif status == "ignored":
                                    ignored_recent[eid] = (
                                        time.time() + ignored_recheck_seconds
                                    )
                            finally:
                                in_progress.discard(eid)

                            if status == "retry":
                                log(
                                    f"Entry {eid}: comment chua xac minh duoc; "
                                    "dung batch hien tai de tranh tao them pending/API cost."
                                )
                                break

                        time.sleep(interval)

                    except KeyboardInterrupt:
                        raise

                    except RuntimeError as exc:
                        code = str(exc)

                        if code == "login_required":
                            log("Phien LMS da het/bi vang. Dang xac minh va dang nhap lai theo config...")
                            user_id, display_name = ensure_login(page, cfg)
                            new_account_key = f"{expected_username}:{user_id}"
                            blog_url = f"{base_url}/blog/index.php?userid={user_id}"
                            if new_account_key != account_key:
                                log(
                                    f"Tai khoan sau reconnect thay doi {account_key} -> "
                                    f"{new_account_key}; nap state rieng."
                                )
                                account_key = new_account_key
                                processed = load_processed(account_key)
                                prune_pending_comments(account_key, processed)
                            scan_floor_by_account[account_key] = ensure_account_scan_floor(
                                page,
                                blog_url,
                                account_key,
                                processed,
                                only_new,
                            )
                            log(
                                f"Dang nhap lai dung tai khoan: "
                                f"{display_name or expected_username} | userid={user_id}"
                            )
                            continue

                        if (
                            code in ("browser_restart_required", "process_restart_required")
                            or code.startswith(("state_corrupt:", "ai_permanent:"))
                        ):
                            raise

                        log(f"Loi tac vu: {exc}")
                        log(f"Cho {interval}s roi thu lai...")
                        time.sleep(interval)

                    except Exception as exc:
                        log(
                            f"Mat ket noi/loi browser tam thoi: "
                            f"{type(exc).__name__}: {exc}"
                        )
                        raise RuntimeError("browser_restart_required") from exc

            except KeyboardInterrupt:
                log("Da dung bot.")
                clear_runtime_pid()
                release_single_instance(instance_lock)
                return 0

            except Exception as exc:
                error_code = str(exc)
                if error_code == "process_restart_required":
                    log("Supervisor se khoi dong lai toan bo bot de giai phong tai nguyen.")
                    write_health(
                        "restarting_process",
                        memory_mb=get_process_memory_mb(),
                        browser_children_mb=find_browser_memory_mb(os.getpid()),
                    )
                    clear_runtime_pid()
                    release_single_instance(instance_lock)
                    return 75
                if error_code.startswith("state_corrupt:"):
                    log(
                        "LOI STATE NGHIEM TRONG: file runtime bi hong va khong co "
                        f"backup hop le ({error_code}). Bot dung de tranh comment trung."
                    )
                    write_health("state_error", error=error_code)
                    clear_runtime_pid()
                    release_single_instance(instance_lock)
                    return 3
                if error_code.startswith("ai_permanent:"):
                    log(
                        "LOI AI KHONG NEN RETRY: kiem tra API key/quyen/model. "
                        f"Chi tiet: {error_code}"
                    )
                    write_health("ai_config_error", error=error_code)
                    clear_runtime_pid()
                    release_single_instance(instance_lock)
                    return 3
                if error_code == "browser_restart_required":
                    log("Can khoi tao lai phien browser.")
                else:
                    log(
                        f"Phien LMS/browser gap loi: "
                        f"{type(exc).__name__}: {exc}"
                    )

            finally:
                if context is not None:
                    try:
                        write_health("browser_closed")
                        context.close()
                    except Exception:
                        pass

            log(f"Host van dang chay. Cho {reconnect_delay}s roi ket noi lai...")
            time.sleep(reconnect_delay)

if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        message = str(exc)
        log(f"Bot khong khoi dong duoc: {type(exc).__name__}: {message}")
        fatal_config = (
            "config.json" in message.lower()
            or "config khong hop le" in message.lower()
        )
        sys.exit(3 if fatal_config else 1)
